# 書き出した pptx を deck.json と突き合わせる: スライド数・各スライドの文字・表・グラフ・画像・ノート。
# 使い方: python tools/check-pptx.py <pptx> <deck.json> [画像の保存先]
# 数式は画像になるので、取り出して保存する(目で見て確かめる)。
import json
import os
import sys

from pptx import Presentation
from pptx.util import Emu

pptx_path, deck_path = sys.argv[1], sys.argv[2]
media_dir = sys.argv[3] if len(sys.argv) > 3 else None
deck = json.load(open(deck_path, encoding="utf-8"))
prs = Presentation(pptx_path)
problems = []

# スライドは 1280x720 px、ノートは 768x1024 px(96px = 1inch)
want_w, want_h = (8.0, 10.67) if deck.get("format") == "note" else (13.33, 7.5)
if round(Emu(prs.slide_width).inches, 2) != want_w or round(Emu(prs.slide_height).inches, 2) != want_h:
    problems.append(f"スライドの大きさ {Emu(prs.slide_width).inches} x {Emu(prs.slide_height).inches}")
if len(prs.slides) != len(deck["slides"]):
    problems.append(f"スライド数 {len(prs.slides)} != {len(deck['slides'])}")

counts = {"text": 0, "table": 0, "chart": 0, "picture": 0}
for i, (s, ds) in enumerate(zip(prs.slides, deck["slides"]), 1):
    texts = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    for el in ds["elements"]:
        if el["type"] == "text":
            want = "".join(r["text"] for p in el["paragraphs"] for r in p["runs"]).replace("\n", "")
            if want.strip() and want.replace(" ", "")[:12] not in texts.replace("\n", "").replace("\x0b", "").replace(" ", ""):
                problems.append(f"スライド {i}: 文字「{want[:20]}」が見つからない")
            counts["text"] += 1
    tables = [sh for sh in s.shapes if sh.has_table]
    charts = [sh for sh in s.shapes if sh.has_chart]
    pictures = [sh for sh in s.shapes if sh.shape_type == 13]
    need = {k: sum(1 for e in ds["elements"] if e["type"] == k) for k in ("table", "chart", "image", "math")}
    if len(tables) != need["table"]: problems.append(f"スライド {i}: 表 {len(tables)} != {need['table']}")
    if len(charts) != need["chart"]: problems.append(f"スライド {i}: グラフ {len(charts)} != {need['chart']}")
    if len(pictures) != need["image"] + need["math"]: problems.append(f"スライド {i}: 画像 {len(pictures)} != {need['image'] + need['math']}")
    for t, el in zip(tables, [e for e in ds["elements"] if e["type"] == "table"]):
        got = [[c.text for c in r.cells] for r in t.table.rows]
        if got != [[c["text"] for c in r] for r in el["rows"]]: problems.append(f"スライド {i}: 表の中身が違う")
    for c, el in zip(charts, [e for e in ds["elements"] if e["type"] == "chart"]):
        plot = c.chart.plots[0]
        if list(plot.categories) != el["chart"]["labels"]: problems.append(f"スライド {i}: グラフのラベルが違う")
    notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    if ds.get("notes", "").strip() != notes.strip(): problems.append(f"スライド {i}: ノートが違う")
    counts["table"] += len(tables); counts["chart"] += len(charts); counts["picture"] += len(pictures)
    if media_dir:
        os.makedirs(media_dir, exist_ok=True)
        for k, p in enumerate(pictures):
            open(os.path.join(media_dir, f"s{i}-{k}.{p.image.ext}"), "wb").write(p.image.blob)

print(f"  pptx: {len(prs.slides)} slides, text {counts['text']}, table {counts['table']}, chart {counts['chart']}, picture {counts['picture']}")
if problems:
    print("\n".join("  NG " + p for p in problems))
    sys.exit(1)
print("  pptx の中身は deck.json と一致")
